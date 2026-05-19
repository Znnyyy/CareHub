from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.http import FileResponse
from django.utils import timezone
from django.db import transaction
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import os
from .models import Report
from .serializers import ReportSerializer
from apps.medical_records.models import MedicalRecord
from apps.medicine.models import Medicine
from apps.patients.models import Patient

class ReportViewSet(viewsets.ModelViewSet):
    queryset = Report.objects.all().order_by('-created_at')
    serializer_class = ReportSerializer
    permission_classes = [permissions.IsAuthenticated]

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        data = request.data.copy()
        data['generated_by'] = request.user.id
        serializer = ReportSerializer(data=data)
        if serializer.is_valid():
            report = serializer.save(generated_by=request.user)
            # Generate file berdasarkan format
            if report.format == 'xlsx':
                file_path = self.generate_xlsx(report)
            else:
                file_path = self.generate_pdf(report)
            report.file = file_path
            report.save()
            return Response(ReportSerializer(report).data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def get_queryset_by_type(self, report):
        if report.report_type == 'visit':
            return MedicalRecord.objects.filter(
                visit_date__range=[report.date_from, report.date_to]
            )
        elif report.report_type == 'medicine':
            return Medicine.objects.filter(is_active=True)
        elif report.report_type == 'patient':
            return Patient.objects.filter(is_active=True)
        return []

    def generate_xlsx(self, report):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = report.title

        # Header style
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        header_alignment = Alignment(horizontal='center')

        if report.report_type == 'visit':
            headers = ['No', 'No. Rekam Medis', 'Nama Pasien', 'Dokter', 'Tanggal', 'Diagnosis', 'Status']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            records = self.get_queryset_by_type(report)
            for row, record in enumerate(records, 2):
                ws.cell(row=row, column=1, value=row-1)
                ws.cell(row=row, column=2, value=record.record_number)
                ws.cell(row=row, column=3, value=record.patient.full_name)
                ws.cell(row=row, column=4, value=f"Dr. {record.doctor.full_name}")
                ws.cell(row=row, column=5, value=str(record.visit_date))
                ws.cell(row=row, column=6, value=record.diagnosis)
                ws.cell(row=row, column=7, value=record.status)

        elif report.report_type == 'medicine':
            headers = ['No', 'Kode', 'Nama Obat', 'Kategori', 'Satuan', 'Stok', 'Stok Min', 'Harga']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            medicines = self.get_queryset_by_type(report)
            for row, medicine in enumerate(medicines, 2):
                ws.cell(row=row, column=1, value=row-1)
                ws.cell(row=row, column=2, value=medicine.code)
                ws.cell(row=row, column=3, value=medicine.name)
                ws.cell(row=row, column=4, value=medicine.category)
                ws.cell(row=row, column=5, value=medicine.unit)
                ws.cell(row=row, column=6, value=medicine.stock)
                ws.cell(row=row, column=7, value=medicine.min_stock)
                ws.cell(row=row, column=8, value=float(medicine.price))

        elif report.report_type == 'patient':
            headers = ['No', 'NIK', 'Nama Lengkap', 'Gender', 'Tanggal Lahir', 'Goldar', 'No. HP']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            patients = self.get_queryset_by_type(report)
            for row, patient in enumerate(patients, 2):
                ws.cell(row=row, column=1, value=row-1)
                ws.cell(row=row, column=2, value=patient.nik)
                ws.cell(row=row, column=3, value=patient.full_name)
                ws.cell(row=row, column=4, value=patient.gender)
                ws.cell(row=row, column=5, value=str(patient.birth_date))
                ws.cell(row=row, column=6, value=patient.blood_type)
                ws.cell(row=row, column=7, value=patient.phone)

        # Auto column width
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_length + 4

        file_name = f"reports/report_{report.id}_{timezone.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        file_path = f"media/{file_name}"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        wb.save(file_path)
        return file_name

    def generate_pdf(self, report):
        file_name = f"reports/report_{report.id}_{timezone.now().strftime('%Y%m%d%H%M%S')}.pdf"
        file_path = f"media/{file_name}"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)

        doc = SimpleDocTemplate(file_path, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []

        # Title
        elements.append(Paragraph(f"<b>{report.title}</b>", styles['Title']))
        elements.append(Paragraph(f"Periode: {report.date_from} s/d {report.date_to}", styles['Normal']))
        elements.append(Spacer(1, 20))

        if report.report_type == 'visit':
            data = [['No', 'No. RM', 'Pasien', 'Dokter', 'Tanggal', 'Status']]
            records = self.get_queryset_by_type(report)
            for i, record in enumerate(records, 1):
                data.append([
                    str(i), record.record_number, record.patient.full_name,
                    f"Dr. {record.doctor.full_name}", str(record.visit_date), record.status
                ])
        elif report.report_type == 'medicine':
            data = [['No', 'Kode', 'Nama Obat', 'Stok', 'Stok Min', 'Harga']]
            medicines = self.get_queryset_by_type(report)
            for i, medicine in enumerate(medicines, 1):
                data.append([
                    str(i), medicine.code, medicine.name,
                    str(medicine.stock), str(medicine.min_stock), f"Rp {medicine.price:,}"
                ])
        elif report.report_type == 'patient':
            data = [['No', 'NIK', 'Nama Lengkap', 'Gender', 'Goldar', 'No. HP']]
            patients = self.get_queryset_by_type(report)
            for i, patient in enumerate(patients, 1):
                data.append([
                    str(i), patient.nik, patient.full_name,
                    patient.gender, patient.blood_type or '-', patient.phone
                ])
        else:
            data = [['Tidak ada data']]

        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ]))
        elements.append(table)
        doc.build(elements)
        return file_name

    @action(detail=True, methods=['get'])
    def download(self, request, pk=None):
        report = self.get_object()
        if not report.file:
            return Response({'error': 'File tidak tersedia'}, status=status.HTTP_404_NOT_FOUND)
        file_path = f"media/{report.file}"
        return FileResponse(open(file_path, 'rb'), as_attachment=True)